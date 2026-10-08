#!/usr/bin/env python3
"""The build-a-schedule skill's helper: drafts a camp activity schedule, checks it, and
writes the printable Excel workbook.

    python schedule.py draft <setup.json> <schedule.json> [--seed N]
    python schedule.py check <schedule.json | workbook.xlsx>
    python schedule.py build <schedule.json> <workbook.xlsx>
    python schedule.py table <schedule.json>

draft  fills every open period it can, fixed periods first, and writes schedule.json.
check  reports any area booked by more groups than it holds in one period (the one rule
       that is always on), then any rule the director asked for that is not met, then
       the open periods. Given a workbook, it reads the Master sheet back from the file.
build  checks first and refuses to write a workbook with a double booking. After writing,
       it opens the file again and checks what is actually in it, Master sheet and area
       sheets both, and deletes the file if anything disagrees.
table  prints one table per day (groups down, periods across) to paste anywhere.

Exit codes: 0 fine, 1 a double booking or a broken name (check refuses, build refuses),
2 the input could not be read.

The setup file (JSON). Only "groups", "days", "periods", "areas" and "activities" are needed:
  {"title": "Summer activity schedule",
   "groups": ["Group 1", ...], "days": ["Day 1", ...], "periods": ["Period 1", ...],
   "areas": [{"name": "Pool", "holds": 1}, {"name": "Dining hall", "holds": "any"}],
   "activities": [{"name": "Swim", "area": "Pool"},
                  {"name": "Archery", "area": "Range", "groups": ["Group 1", "Group 2"]},
                  {"name": "Lunch", "area": "Dining hall"}],
   "rules": {...}}
An area holds one group at a time unless "holds" says a number or "any" (the director
marked it shareable). An activity's "groups" list, when present, limits who may do it.

"rules" holds only what the director asked for. Nothing in it is on by default:
  "fixed": [{"groups": "all" or [...], "days": "all" or [...], "period": "Period 4",
             "activity": "Lunch"}]      the same activity at the same time
  "no_repeat_in_a_day": true            no group does one activity twice in a day
  "times_per_cycle": {"Swim": 2}        each group does it exactly this many times
  "most_at_once": {"Archery": 1}        a staff limit: at most this many groups at once

A schedule file is the setup plus "grid": {group: {day: {period: activity}}}. A missing
or empty period is open.

Python 3.8 or newer. Only build needs openpyxl; check on a JSON file and draft do not.
"""
import json
import os
import random
import sys

OPEN_VALUES = ("", "open", "free")  # read back from a workbook as an open period
CHECKER_SHEET = "_checker"
MASTER = "Master"
SETUP_SHEET = "Setup"


class BadInput(Exception):
    pass


# --------------------------------------------------------------------------
# Reading and checking the setup
# --------------------------------------------------------------------------

def load_json(path):
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError) as e:
        raise BadInput(f"Could not read {os.path.basename(path)}: {e}")


def norm(s):
    return " ".join(str(s).split()).casefold()


def capacity(area):
    """How many groups the area holds in one period; None means any number."""
    h = area.get("holds", 1)
    if isinstance(h, str) and h.strip().lower() in ("any", "shared", "shareable", "unlimited"):
        return None
    try:
        h = int(h)
    except (TypeError, ValueError):
        raise BadInput(f"Area {area.get('name')!r}: 'holds' must be a number or \"any\", not {h!r}")
    if h < 1:
        raise BadInput(f"Area {area.get('name')!r}: 'holds' must be at least 1")
    return h


def validate_setup(s):
    for key in ("groups", "days", "periods", "areas", "activities"):
        if not isinstance(s.get(key), list) or not s[key]:
            raise BadInput(f"The setup needs a non-empty list called {key!r}")
    for key in ("groups", "days", "periods"):
        seen = set()
        for name in s[key]:
            if not isinstance(name, str) or not name.strip():
                raise BadInput(f"Every name in {key!r} must be words, not {name!r}")
            if norm(name) in seen:
                raise BadInput(f"{name!r} appears twice in {key!r}")
            seen.add(norm(name))
    areas = {}
    for a in s["areas"]:
        if not isinstance(a, dict) or not a.get("name"):
            raise BadInput("Every area needs a name")
        if norm(a["name"]) in areas:
            raise BadInput(f"Area {a['name']!r} appears twice")
        areas[norm(a["name"])] = (a["name"], capacity(a))
    acts = {}
    for a in s["activities"]:
        if not isinstance(a, dict) or not a.get("name"):
            raise BadInput("Every activity needs a name")
        if norm(a["name"]) in acts:
            raise BadInput(f"Activity {a['name']!r} appears twice")
        if norm(a.get("area", "")) not in areas:
            raise BadInput(f"Activity {a['name']!r} is in area {a.get('area')!r}, which is not in the areas list")
        if norm(a["name"]) in OPEN_VALUES:
            raise BadInput(f"{a['name']!r} cannot be an activity name; it means an open period")
        groups = a.get("groups")
        if groups is not None:
            known = {norm(g) for g in s["groups"]}
            for g in groups:
                if norm(g) not in known:
                    raise BadInput(f"Activity {a['name']!r} lists group {g!r}, which is not in the groups list")
        acts[norm(a["name"])] = a
    rules = s.get("rules") or {}
    if not isinstance(rules, dict):
        raise BadInput("'rules' must be a set of named rules")
    unknown = set(rules) - {"fixed", "no_repeat_in_a_day", "times_per_cycle", "most_at_once"}
    if unknown:
        raise BadInput(f"Unknown rule(s): {', '.join(sorted(unknown))}. Check those by hand and leave them out of the file.")
    for f in rules.get("fixed", []):
        if norm(f.get("activity", "")) not in acts:
            raise BadInput(f"A fixed period names activity {f.get('activity')!r}, which is not in the activities list")
        if norm(f.get("period", "")) not in {norm(p) for p in s["periods"]}:
            raise BadInput(f"A fixed period names period {f.get('period')!r}, which is not in the periods list")
    for key in ("times_per_cycle", "most_at_once"):
        for name, n in (rules.get(key) or {}).items():
            if norm(name) not in acts:
                raise BadInput(f"Rule {key!r} names activity {name!r}, which is not in the activities list")
            if not isinstance(n, int) or n < 0:
                raise BadInput(f"Rule {key!r} for {name!r} must be a whole number")
    return areas, acts


def expand(spec, names):
    if spec in (None, "all", "*"):
        return list(names)
    by = {norm(n): n for n in names}
    return [by[norm(x)] for x in spec if norm(x) in by]


# --------------------------------------------------------------------------
# The check
# --------------------------------------------------------------------------

def check(s, grid):
    """Returns (hard, rule_misses, open_slots, filled, total). hard holds every double
    booking and every name the setup does not know; rule_misses only rules she asked for."""
    areas, acts = validate_setup(s)
    rules = s.get("rules") or {}
    hard, misses, opens = [], [], []
    gnames = {norm(g): g for g in s["groups"]}
    dnames = {norm(d): d for d in s["days"]}
    pnames = {norm(p): p for p in s["periods"]}
    for g, days in grid.items():
        if norm(g) not in gnames:
            hard.append(f"The schedule has a group called {g!r} that is not in the setup.")
            continue
        for d, periods in days.items():
            if norm(d) not in dnames:
                hard.append(f"{g} has a day called {d!r} that is not in the setup.")
                continue
            for p, act in periods.items():
                if norm(p) not in pnames:
                    hard.append(f"{g}, {d} has a period called {p!r} that is not in the setup.")
                elif norm(act or "") not in OPEN_VALUES and norm(act) not in acts:
                    hard.append(f"{g}, {d}, {p}: {act!r} is not in the activities list, so its area is unknown.")

    def cell(g, d, p):
        v = (((grid.get(g) or {}).get(d) or {}).get(p) or "")
        return "" if norm(v) in OPEN_VALUES else v

    filled = total = 0
    for d in s["days"]:
        for p in s["periods"]:
            in_area, in_act = {}, {}
            for g in s["groups"]:
                total += 1
                v = cell(g, d, p)
                if not v:
                    opens.append((g, d, p))
                    continue
                filled += 1
                a = acts.get(norm(v))
                if not a:
                    continue
                in_area.setdefault(norm(a["area"]), []).append(g)
                in_act.setdefault(norm(v), []).append(g)
            for key, gs in in_area.items():
                name, cap = areas[key]
                if cap is not None and len(gs) > cap:
                    hard.append(f"Double booked: {name}, {d}, {p}: {', '.join(gs)} "
                                f"({len(gs)} groups; it holds {cap}).")
            for name, n in (rules.get("most_at_once") or {}).items():
                gs = in_act.get(norm(name), [])
                if len(gs) > n:
                    misses.append(f"Staff limit: {name}, {d}, {p} has {len(gs)} groups ({', '.join(gs)}); "
                                  f"you said at most {n}.")
    for f in rules.get("fixed", []):
        for g in expand(f.get("groups"), s["groups"]):
            for d in expand(f.get("days"), s["days"]):
                p = pnames[norm(f["period"])]
                if norm(cell(g, d, p)) != norm(f["activity"]):
                    misses.append(f"Fixed period: {g}, {d}, {p} should be {f['activity']} "
                                  f"and is {cell(g, d, p) or 'open'}.")
    for g in s["groups"]:
        for a in acts.values():
            allowed = a.get("groups")
            if allowed is None:
                continue
            if norm(g) not in {norm(x) for x in allowed}:
                for d in s["days"]:
                    for p in s["periods"]:
                        if norm(cell(g, d, p)) == norm(a["name"]):
                            misses.append(f"{g} has {a['name']} on {d}, {p}, but {a['name']} is only for "
                                          f"{', '.join(allowed)}.")
        if rules.get("no_repeat_in_a_day"):
            for d in s["days"]:
                seen = {}
                for p in s["periods"]:
                    v = cell(g, d, p)
                    if v:
                        seen.setdefault(norm(v), []).append(p)
                for v, ps in seen.items():
                    if len(ps) > 1:
                        misses.append(f"Repeat: {g} has {acts[v]['name'] if v in acts else v} twice on {d} "
                                      f"({', '.join(ps)}).")
        for name, n in (rules.get("times_per_cycle") or {}).items():
            got = sum(1 for d in s["days"] for p in s["periods"] if norm(cell(g, d, p)) == norm(name))
            if got != n:
                misses.append(f"Rotation: {g} has {name} {got} time(s) this cycle; you said {n}.")
    return hard, misses, opens, filled, total


def report(s, grid, out=sys.stdout):
    hard, misses, opens, filled, total = check(s, grid)
    doubles = [h for h in hard if h.startswith("Double booked")]
    if not hard:
        out.write("No area is booked by more groups than it holds, in any period.\n")
    else:
        out.write(f"{len(doubles)} double booking(s) and {len(hard) - len(doubles)} other problem(s):\n")
        for h in hard:
            out.write("  " + h + "\n")
    if (s.get("rules") or {}):
        if misses:
            out.write(f"{len(misses)} place(s) where a rule you asked for is not met:\n")
            for m in misses:
                out.write("  " + m + "\n")
        else:
            out.write("Every rule you asked for is met.\n")
    out.write(f"{filled} of {total} periods filled, {len(opens)} open.\n")
    for g, d, p in opens[:40]:
        out.write(f"  Open: {g}, {d}, {p}\n")
    if len(opens) > 40:
        out.write(f"  ...and {len(opens) - 40} more open.\n")
    return 1 if hard else 0


# --------------------------------------------------------------------------
# The draft
# --------------------------------------------------------------------------

def draft(s, seed=None, attempts=60):
    areas, acts = validate_setup(s)
    rules = s.get("rules") or {}
    rng = random.Random(seed)
    fixed_only = {norm(f["activity"]) for f in rules.get("fixed", [])}
    fill_with = [a for k, a in acts.items() if k not in fixed_only or a.get("also_elsewhere")]
    need = {norm(k): v for k, v in (rules.get("times_per_cycle") or {}).items()}
    most = {norm(k): v for k, v in (rules.get("most_at_once") or {}).items()}
    best, best_score = None, None
    for _ in range(attempts):
        grid = {g: {d: {} for d in s["days"]} for g in s["groups"]}
        done = {g: {} for g in s["groups"]}
        for f in rules.get("fixed", []):
            p = next(x for x in s["periods"] if norm(x) == norm(f["period"]))
            act = acts[norm(f["activity"])]["name"]
            for g in expand(f.get("groups"), s["groups"]):
                for d in expand(f.get("days"), s["days"]):
                    grid[g][d][p] = act
                    done[g][norm(act)] = done[g].get(norm(act), 0) + 1
        slots_left = {g: sum(1 for d in s["days"] for p in s["periods"] if p not in grid[g][d])
                      for g in s["groups"]}
        for d in s["days"]:
            for p in s["periods"]:
                used_area, used_act = {}, {}
                for g in s["groups"]:
                    v = grid[g][d].get(p)
                    if v:
                        a = acts[norm(v)]
                        used_area[norm(a["area"])] = used_area.get(norm(a["area"]), 0) + 1
                        used_act[norm(v)] = used_act.get(norm(v), 0) + 1
                order = [g for g in s["groups"] if p not in grid[g][d]]
                rng.shuffle(order)
                # Groups with the most required activities still owed choose first.
                order.sort(key=lambda g: -sum(max(0, n - done[g].get(k, 0)) for k, n in need.items()))
                for g in order:
                    today = {norm(v) for v in grid[g][d].values()}
                    choices = []
                    for a in fill_with:
                        k = norm(a["name"])
                        cap = areas[norm(a["area"])][1]
                        if cap is not None and used_area.get(norm(a["area"]), 0) >= cap:
                            continue
                        if k in most and used_act.get(k, 0) >= most[k]:
                            continue
                        if a.get("groups") is not None and norm(g) not in {norm(x) for x in a["groups"]}:
                            continue
                        if rules.get("no_repeat_in_a_day") and k in today:
                            continue
                        if k in need and done[g].get(k, 0) >= need[k]:
                            continue
                        owed = max(0, need.get(k, 0) - done[g].get(k, 0))
                        urgency = owed / max(1, slots_left[g])
                        # Without a rotation rule this only spreads activities so the draft
                        # reads sensibly; the check never enforces it.
                        choices.append((-urgency, done[g].get(k, 0), rng.random(), a))
                    slots_left[g] -= 1
                    if not choices:
                        continue
                    a = min(choices, key=lambda c: c[:3])[3]
                    grid[g][d][p] = a["name"]
                    done[g][norm(a["name"])] = done[g].get(norm(a["name"]), 0) + 1
                    used_area[norm(a["area"])] = used_area.get(norm(a["area"]), 0) + 1
                    used_act[norm(a["name"])] = used_act.get(norm(a["name"]), 0) + 1
        repair(s, grid, areas, acts, fill_with, rules, rng)
        hard, misses, opens, _, _ = check(s, grid)
        score = (len(hard), len(misses), len(opens))
        if best_score is None or score < best_score:
            best, best_score = grid, score
        if score == (0, 0, 0):
            break
    return best


def can_take(s, grid, g, d, p, a, areas, rules, ignore=None):
    """Whether group g may do activity a at (d, p) under the hard rule and the rules she
    asked for. ignore is a group whose current activity at (d, p) is about to move."""
    k = norm(a["name"])
    if a.get("groups") is not None and norm(g) not in {norm(x) for x in a["groups"]}:
        return False
    others = [x for x in s["groups"] if x not in (g, ignore) and grid[x][d].get(p)]
    cap = areas[norm(a["area"])][1]
    if cap is not None:
        inside = [x for x in others if norm(_area(grid[x][d][p], s)) == norm(a["area"])]
        if len(inside) >= cap:
            return False
    most = {norm(n): v for n, v in (rules.get("most_at_once") or {}).items()}
    if k in most and sum(1 for x in others if norm(grid[x][d][p]) == k) >= most[k]:
        return False
    if rules.get("no_repeat_in_a_day") and any(norm(v) == k for q, v in grid[g][d].items() if q != p):
        return False
    need = {norm(n): v for n, v in (rules.get("times_per_cycle") or {}).items()}
    if k in need:
        have = sum(1 for dd in s["days"] for q, v in grid[g][dd].items() if norm(v) == k and (dd, q) != (d, p))
        if have >= need[k]:
            return False
    return True


def _area(activity, s):
    for a in s["activities"]:
        if norm(a["name"]) == norm(activity):
            return a["area"]
    return ""


def repair(s, grid, areas, acts, fill_with, rules, rng):
    """Fills open periods a straight pass left behind by moving one other group in the
    same period to a different activity. Never touches a fixed period."""
    need = {norm(n) for n in (rules.get("times_per_cycle") or {})}
    fixed = set()
    for f in rules.get("fixed", []):
        p = next(x for x in s["periods"] if norm(x) == norm(f["period"]))
        for g in expand(f.get("groups"), s["groups"]):
            for d in expand(f.get("days"), s["days"]):
                fixed.add((g, d, p))
    for d in s["days"]:
        for p in s["periods"]:
            for g in s["groups"]:
                if grid[g][d].get(p):
                    continue
                choices = list(fill_with)
                rng.shuffle(choices)
                placed = False
                for a in choices:
                    if can_take(s, grid, g, d, p, a, areas, rules):
                        grid[g][d][p] = a["name"]
                        placed = True
                        break
                if placed:
                    continue
                for a in choices:
                    for h in s["groups"]:
                        cur = grid[h][d].get(p)
                        if h == g or not cur or (h, d, p) in fixed or norm(cur) in need:
                            continue
                        if not can_take(s, grid, g, d, p, a, areas, rules, ignore=h):
                            continue
                        if norm(_area(cur, s)) != norm(a["area"]) and norm(cur) != norm(a["name"]):
                            continue
                        old = grid[h][d].pop(p)
                        grid[g][d][p] = a["name"]
                        moved = next((b for b in choices if norm(b["name"]) != norm(old)
                                      and can_take(s, grid, h, d, p, b, areas, rules)), None)
                        if moved:
                            grid[h][d][p] = moved["name"]
                            placed = True
                            break
                        del grid[g][d][p]
                        grid[h][d][p] = old
                    if placed:
                        break


# --------------------------------------------------------------------------
# The workbook
# --------------------------------------------------------------------------

def sheet_title(name, taken):
    bad = '[]:*?/\\'
    t = "".join("-" if c in bad else c for c in str(name)).strip().strip("'") or "Sheet"
    t = t[:31]
    base, i = t, 2
    while t.casefold() in taken:
        suffix = f" ({i})"
        t = base[:31 - len(suffix)] + suffix
        i += 1
    taken.add(t.casefold())
    return t


def build(s, grid, path):
    hard, _, _, _, _ = check(s, grid)
    if hard:
        sys.stdout.write("Refused: the schedule has problems, so no workbook was written.\n")
        report(s, grid)
        return 1
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.stdout.write("openpyxl is not installed, so the workbook cannot be made. "
                         "Try: python -m pip install openpyxl\n")
        return 2
    areas, acts = validate_setup(s)
    title = s.get("title") or "Activity schedule"
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = PatternFill("solid", fgColor="DDE6EE")
    open_fill = PatternFill("solid", fgColor="FFF2CC")
    wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)
    bold = Font(bold=True)

    def val(g, d, p):
        v = (((grid.get(g) or {}).get(d) or {}).get(p) or "")
        return "" if norm(v) in OPEN_VALUES else acts[norm(v)]["name"]

    def setup_print(ws, landscape, rows_repeat):
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.gridLines = False
        ws.print_options.horizontalCentered = True
        ws.print_title_rows = rows_repeat
        ws.oddFooter.center.text = title + " - page &P of &N"

    wb = Workbook()
    taken = {MASTER.casefold(), SETUP_SHEET.casefold(), CHECKER_SHEET.casefold()}
    ws = wb.active
    ws.title = MASTER
    ws["A1"] = title + ": every group"
    ws["A1"].font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value="Group").font = bold
    ws.cell(row=3, column=1, value="")
    np_ = len(s["periods"])
    for di, d in enumerate(s["days"]):
        c0 = 2 + di * np_
        ws.cell(row=2, column=c0, value=d)
        if np_ > 1:
            ws.merge_cells(start_row=2, start_column=c0, end_row=2, end_column=c0 + np_ - 1)
        for pi, p in enumerate(s["periods"]):
            c = ws.cell(row=3, column=c0 + pi, value=p)
            c.font, c.fill, c.alignment, c.border = bold, head, wrap, box
        hc = ws.cell(row=2, column=c0)
        hc.font, hc.fill, hc.alignment = bold, head, wrap
    for gi, g in enumerate(s["groups"]):
        r = 4 + gi
        c = ws.cell(row=r, column=1, value=g)
        c.font, c.border = bold, box
        for di, d in enumerate(s["days"]):
            for pi, p in enumerate(s["periods"]):
                c = ws.cell(row=r, column=2 + di * np_ + pi, value=val(g, d, p))
                c.alignment, c.border = wrap, box
                if not c.value:
                    c.fill = open_fill
    ws.column_dimensions["A"].width = 16
    for col in range(2, 2 + len(s["days"]) * np_):
        ws.column_dimensions[get_column_letter(col)].width = 12
    ws.freeze_panes = "B4"
    setup_print(ws, True, "2:3")

    def grid_sheet(name, heading, note, value_of):
        sh = wb.create_sheet(sheet_title(name, taken))
        sh["A1"] = heading
        sh["A1"].font = Font(bold=True, size=14)
        if note:
            sh["A2"] = note
        sh.cell(row=3, column=1, value="").border = box
        for di, d in enumerate(s["days"]):
            c = sh.cell(row=3, column=2 + di, value=d)
            c.font, c.fill, c.alignment, c.border = bold, head, wrap, box
        for pi, p in enumerate(s["periods"]):
            c = sh.cell(row=4 + pi, column=1, value=p)
            c.font, c.fill, c.alignment, c.border = bold, head, wrap, box
            sh.row_dimensions[4 + pi].height = 36
            for di, d in enumerate(s["days"]):
                c = sh.cell(row=4 + pi, column=2 + di, value=value_of(d, p))
                c.alignment, c.border = wrap, box
        sh.column_dimensions["A"].width = 14
        for col in range(2, 2 + len(s["days"])):
            sh.column_dimensions[get_column_letter(col)].width = 18
        setup_print(sh, len(s["days"]) > 5, "3:3")
        return sh

    for g in s["groups"]:
        def gv(d, p, g=g):
            v = val(g, d, p)
            if not v:
                return ""
            area = acts[norm(v)]["area"]
            return v if norm(area) == norm(v) else f"{v}\n({area})"
        grid_sheet(g, f"{title}: {g}", "", gv)
    area_sheets = {}
    for key, (aname, cap) in areas.items():
        note = "One group at a time." if cap == 1 else (
            "Shared: any number of groups at once." if cap is None else f"Holds {cap} groups at once.")

        def av(d, p, key=key):
            return ", ".join(g for g in s["groups"] if val(g, d, p) and norm(acts[norm(val(g, d, p))]["area"]) == key)
        area_sheets[key] = grid_sheet("Area - " + aname, f"{title}: {aname}", note, av).title

    st = wb.create_sheet(SETUP_SHEET)
    st["A1"] = "Areas and activities"
    st["A1"].font = Font(bold=True, size=14)
    st.append([])
    st.append(["Area", "Groups at once"])
    for aname, cap in areas.values():
        st.append([aname, "any" if cap is None else cap])
    st.append([])
    st.append(["Activity", "Area", "Only for"])
    for a in acts.values():
        st.append([a["name"], a["area"], ", ".join(a.get("groups") or [])])
    st.append([])
    st.append(["Rules you asked for"])
    for line in describe_rules(s):
        st.append([line])
    st.column_dimensions["A"].width = 40
    st.column_dimensions["B"].width = 18
    st.column_dimensions["C"].width = 30
    setup_print(st, False, None)

    ck = wb.create_sheet(CHECKER_SHEET)
    ck["A1"] = "Used to re-check this workbook. Please leave this sheet as it is."
    setup = {k: v for k, v in s.items() if k != "grid"}
    setup["area_sheets"] = {areas[k][0]: t for k, t in area_sheets.items()}
    ck["A2"] = json.dumps(setup)
    ck.sheet_state = "hidden"

    tmp = path + ".tmp.xlsx"
    wb.save(tmp)
    problems = verify_workbook(tmp, s, grid)
    if problems:
        os.remove(tmp)
        sys.stdout.write("Refused: the written workbook did not match the schedule, so it was deleted:\n")
        for p in problems:
            sys.stdout.write("  " + p + "\n")
        return 1
    os.replace(tmp, path)
    sys.stdout.write(f"Wrote {os.path.basename(path)}: Master, {len(s['groups'])} group sheets, "
                     f"{len(areas)} area sheets and Setup. Re-read from the file: no area is booked "
                     f"by more groups than it holds.\n")
    return 0


def describe_rules(s):
    r = s.get("rules") or {}
    out = []
    for f in r.get("fixed", []):
        who = "every group" if f.get("groups") in (None, "all", "*") else ", ".join(f["groups"])
        when = "every day" if f.get("days") in (None, "all", "*") else ", ".join(f["days"])
        out.append(f"{f['activity']} is fixed at {f['period']} for {who}, {when}.")
    if r.get("no_repeat_in_a_day"):
        out.append("No group does the same activity twice in one day.")
    for k, n in (r.get("times_per_cycle") or {}).items():
        out.append(f"Each group does {k} {n} time(s) per cycle.")
    for k, n in (r.get("most_at_once") or {}).items():
        out.append(f"At most {n} group(s) do {k} at once.")
    return out or ["None. The only rule is that no area holds more groups than it can."]


def read_workbook(path):
    """Reads the setup and the grid back out of a workbook this script wrote, using the
    Master sheet's cells as the schedule (so edits made in Excel are what gets checked)."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise BadInput("openpyxl is not installed, so the workbook cannot be read. Try: python -m pip install openpyxl")
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as e:
        raise BadInput(f"Could not open {os.path.basename(path)}: {e}")
    if CHECKER_SHEET not in wb.sheetnames or MASTER not in wb.sheetnames:
        raise BadInput("This workbook was not made by this skill (no Master sheet or no checker sheet).")
    try:
        s = json.loads(wb[CHECKER_SHEET]["A2"].value)
    except (TypeError, ValueError):
        raise BadInput("The checker sheet in this workbook was changed, so it cannot be re-checked.")
    ws = wb[MASTER]
    np_ = len(s["periods"])
    grid = {}
    rows = {norm(ws.cell(row=r, column=1).value or ""): r for r in range(4, ws.max_row + 1)}
    for g in s["groups"]:
        r = rows.get(norm(g))
        if r is None:
            raise BadInput(f"The Master sheet has no row for {g}.")
        grid[g] = {}
        for di, d in enumerate(s["days"]):
            grid[g][d] = {}
            for pi, p in enumerate(s["periods"]):
                v = ws.cell(row=r, column=2 + di * np_ + pi).value
                v = "" if v is None else str(v).strip()
                if v:
                    grid[g][d][p] = v
    return s, grid, wb


def verify_workbook(path, s, grid):
    """Independent of the in-memory schedule: re-reads the file, runs the check on what is
    in it, and compares every area sheet with the Master sheet."""
    problems = []
    s2, grid2, wb = read_workbook(path)
    hard, _, _, _, _ = check(s2, grid2)
    problems += hard
    for g in s["groups"]:
        for d in s["days"]:
            for p in s["periods"]:
                a = norm(((grid.get(g) or {}).get(d) or {}).get(p) or "")
                b = norm(grid2[g][d].get(p, ""))
                if (a if a not in OPEN_VALUES else "") != (b if b not in OPEN_VALUES else ""):
                    problems.append(f"Master sheet differs from the schedule at {g}, {d}, {p}.")
    areas, acts = validate_setup(s2)
    for aname, title in s2.get("area_sheets", {}).items():
        sh = wb[title]
        cap = areas[norm(aname)][1]
        for pi, p in enumerate(s2["periods"]):
            for di, d in enumerate(s2["days"]):
                text = sh.cell(row=4 + pi, column=2 + di).value or ""
                listed = [x.strip() for x in str(text).split(",") if x.strip()]
                want = [g for g in s2["groups"] if grid2[g][d].get(p) and
                        norm(acts[norm(grid2[g][d][p])]["area"]) == norm(aname)]
                if listed != want:
                    problems.append(f"Area sheet {title} differs from the Master sheet at {d}, {p}.")
                if cap is not None and len(listed) > cap:
                    problems.append(f"Area sheet {title} shows {len(listed)} groups at {d}, {p}; it holds {cap}.")
    return problems


def table(s, grid, out=sys.stdout):
    for d in s["days"]:
        out.write(f"\n{d}\n\n| Group | " + " | ".join(s["periods"]) + " |\n")
        out.write("|---" * (len(s["periods"]) + 1) + "|\n")
        for g in s["groups"]:
            cells = []
            for p in s["periods"]:
                v = ((grid.get(g) or {}).get(d) or {}).get(p) or ""
                cells.append("Open" if norm(v) in OPEN_VALUES else v)
            out.write(f"| {g} | " + " | ".join(cells) + " |\n")
    return 0


def main(argv):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if len(argv) < 2 or argv[0] not in ("draft", "check", "build", "table"):
        sys.stdout.write(__doc__)
        return 2
    cmd = argv[0]
    try:
        if cmd == "check" and argv[1].lower().endswith(".xlsx"):
            s, grid, _ = read_workbook(argv[1])
            return report(s, grid)
        s = load_json(argv[1])
        if not isinstance(s, dict):
            raise BadInput("The file must hold one setup, with groups, days, periods, areas and activities.")
        grid = s.get("grid") or {}
        if cmd == "draft":
            if len(argv) < 3:
                raise BadInput("draft needs a place to write the schedule: draft <setup.json> <schedule.json>")
            seed = None
            if "--seed" in argv:
                seed = int(argv[argv.index("--seed") + 1])
            s = dict(s, grid=draft(s, seed=seed))
            with open(argv[2], "w", encoding="utf-8") as f:
                json.dump(s, f, indent=1)
            sys.stdout.write(f"Draft written to {os.path.basename(argv[2])}.\n")
            report(s, s["grid"])
            return 0
        if cmd == "check":
            return report(s, grid)
        if cmd == "table":
            validate_setup(s)
            return table(s, grid)
        if len(argv) < 3:
            raise BadInput("build needs a file name for the workbook: build <schedule.json> <workbook.xlsx>")
        return build(s, grid, argv[2])
    except BadInput as e:
        sys.stdout.write(f"Could not do that: {e}\n")
        return 2


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
