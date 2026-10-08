"""Tests the build-a-schedule skill's script on an invented camp: 18 groups, a five-day
cycle, six periods a day. Runs with no paid evals:

    python tests/schedule_test.py

Needs openpyxl. Every check that says "no double booking" is paired with a red control
that plants one and must be caught, so a green here could have gone red.
"""
import json
import os
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPT = os.path.join(HERE, "plugins", "cg-ember", "skills", "build-a-schedule", "scripts", "schedule.py")
fails = []


def ok(cond, label):
    print(("ok   " if cond else "WRONG ") + label)
    if not cond:
        fails.append(label)


def run(*args):
    p = subprocess.run([sys.executable, SCRIPT] + list(args), capture_output=True, text=True,
                       encoding="utf-8", timeout=300)
    return p.returncode, p.stdout


GROUPS = [f"Group {i}" for i in range(1, 19)]
DAYS = [f"Day {i}" for i in range(1, 6)]
PERIODS = [f"Period {i}" for i in range(1, 7)]
SETUP = {
    "title": "Pine Hollow test schedule",
    "groups": GROUPS, "days": DAYS, "periods": PERIODS,
    "areas": [
        {"name": "Pool", "holds": 2}, {"name": "Gaga pit"}, {"name": "Basketball court", "holds": 2},
        {"name": "Field", "holds": 3}, {"name": "Arts room", "holds": 2}, {"name": "Archery range"},
        {"name": "Nature trail", "holds": 2}, {"name": "Music room"}, {"name": "Gym", "holds": 2},
        {"name": "Tennis courts", "holds": 2}, {"name": "Playground", "holds": 3},
        {"name": "Dining hall", "holds": "any"},
    ],
    "activities": [
        {"name": "Swim", "area": "Pool"}, {"name": "Gaga", "area": "Gaga pit"},
        {"name": "Basketball", "area": "Basketball court"}, {"name": "Soccer", "area": "Field"},
        {"name": "Arts and crafts", "area": "Arts room"},
        {"name": "Archery", "area": "Archery range", "groups": GROUPS[9:]},
        {"name": "Nature", "area": "Nature trail"}, {"name": "Music", "area": "Music room"},
        {"name": "Gymnastics", "area": "Gym"}, {"name": "Tennis", "area": "Tennis courts"},
        {"name": "Free play", "area": "Playground"}, {"name": "Lunch", "area": "Dining hall"},
    ],
    "rules": {
        "fixed": [{"groups": "all", "days": "all", "period": "Period 4", "activity": "Lunch"}],
        "no_repeat_in_a_day": True,
        "times_per_cycle": {"Swim": 2},
        "most_at_once": {"Soccer": 2},
    },
}


def main():
    t = tempfile.mkdtemp()
    setup = os.path.join(t, "setup.json")
    sched = os.path.join(t, "schedule.json")
    book = os.path.join(t, "schedule.xlsx")
    json.dump(SETUP, open(setup, "w", encoding="utf-8"))

    code, out = run("draft", setup, sched, "--seed", "7")
    ok(code == 0, "draft runs")
    s = json.load(open(sched, encoding="utf-8"))
    cells = [s["grid"][g][d].get(p, "") for g in GROUPS for d in DAYS for p in PERIODS]
    ok(len(cells) == 18 * 5 * 6, f"the grid has 540 periods (has {len(cells)})")
    ok(all(cells), f"every period is filled ({sum(1 for c in cells if not c)} open)")
    ok("No area is booked by more groups than it holds" in out, "the draft's own check finds no double booking")
    ok("Every rule you asked for is met" in out, "the draft meets the rules asked for")

    # An independent count, not through the script's check: every area, every period.
    holds = {a["name"]: a.get("holds", 1) for a in SETUP["areas"]}
    area_of = {a["name"]: a["area"] for a in SETUP["activities"]}
    worst = 0
    for d in DAYS:
        for p in PERIODS:
            count = {}
            for g in GROUPS:
                a = area_of[s["grid"][g][d][p]]
                count[a] = count.get(a, 0) + 1
            for a, n in count.items():
                if holds[a] != "any":
                    worst = max(worst, n - holds[a])
    ok(worst <= 0, "counted separately, no area holds more groups than it can")
    lunch = all(s["grid"][g][d]["Period 4"] == "Lunch" for g in GROUPS for d in DAYS)
    ok(lunch, "Lunch sits at Period 4 for every group, every day (fixed, as asked)")
    ok(all(s["grid"][g][d][p] != "Archery" for g in GROUPS[:9] for d in DAYS for p in PERIODS),
       "Archery only goes to the groups allowed it")

    code, out = run("build", sched, book)
    ok(code == 0 and os.path.exists(book), "build writes the workbook")
    from openpyxl import load_workbook
    wb = load_workbook(book)
    names = wb.sheetnames
    ok(names[0] == "Master", "the Master sheet comes first")
    ok(all(g in names for g in GROUPS), "one sheet per group (18)")
    ok(sum(1 for n in names if n.startswith("Area - ")) == len(SETUP["areas"]), "one sheet per area (12)")
    ok("Setup" in names and wb["_checker"].sheet_state == "hidden", "Setup sheet visible, checker sheet hidden")
    ok(wb["Master"].page_setup.orientation == "landscape" and wb["Master"].sheet_properties.pageSetUpPr.fitToPage,
       "Master prints landscape, fit to the page width")
    code, out = run("check", book)
    ok(code == 0 and "No area is booked" in out, "checking the workbook itself finds no double booking")

    # Red control 1: a double booking in the schedule file is caught, and build refuses.
    bad = json.loads(json.dumps(s))
    bad["grid"]["Group 1"]["Day 2"]["Period 1"] = "Gaga"
    bad["grid"]["Group 2"]["Day 2"]["Period 1"] = "Gaga"
    badf = os.path.join(t, "bad.json")
    json.dump(bad, open(badf, "w", encoding="utf-8"))
    code, out = run("check", badf)
    ok(code == 1 and "Double booked: Gaga pit, Day 2, Period 1" in out, "red control: a planted double booking is caught")
    badbook = os.path.join(t, "bad.xlsx")
    code, out = run("build", badf, badbook)
    ok(code == 1 and not os.path.exists(badbook), "red control: build refuses and writes no file")

    # Red control 2: an edit made in the workbook itself is caught when the workbook is checked.
    wb = load_workbook(book)
    ws = wb["Master"]
    ws.cell(row=4, column=2).value = "Music"   # Group 1, Day 1, Period 1
    ws.cell(row=5, column=2).value = "Music"   # Group 2, Day 1, Period 1
    edited = os.path.join(t, "edited.xlsx")
    wb.save(edited)
    code, out = run("check", edited)
    ok(code == 1 and "Double booked: Music room, Day 1, Period 1" in out,
       "red control: a double booking typed into the workbook is caught")

    # A shareable area is not a double booking, and a numbered one is held to its number.
    share = json.loads(json.dumps(s))
    for g in GROUPS[:3]:
        share["grid"][g]["Day 3"]["Period 2"] = "Soccer"
    for g in GROUPS[3:]:
        if share["grid"][g]["Day 3"]["Period 2"] == "Soccer":
            share["grid"][g]["Day 3"]["Period 2"] = ""
    sf = os.path.join(t, "share.json")
    json.dump(share, open(sf, "w", encoding="utf-8"))
    code, out = run("check", sf)
    ok("Double booked: Field" not in out, "three groups on the Field (holds 3) is not a double booking")
    ok("Staff limit: Soccer, Day 3, Period 2" in out and code == 0,
       "the staff limit she asked for is reported, and does not block (it is not the hard rule)")
    for g in GROUPS[3:5]:
        share["grid"][g]["Day 3"]["Period 2"] = "Soccer"
    json.dump(share, open(sf, "w", encoding="utf-8"))
    code, out = run("check", sf)
    ok(code == 1 and "Double booked: Field, Day 3, Period 2" in out, "five groups on the Field (holds 3) is caught")

    # With no rules asked for, nothing but the double booking rule is checked.
    plain = dict(SETUP)
    plain.pop("rules")
    pf = os.path.join(t, "plain.json")
    ps = os.path.join(t, "plain-schedule.json")
    json.dump(plain, open(pf, "w", encoding="utf-8"))
    code, out = run("draft", pf, ps, "--seed", "3")
    ok(code == 0 and "rule you asked for" not in out and "No area is booked" in out,
       "with no rules asked for, only the double booking rule is reported")

    # A name the setup does not know is refused rather than guessed.
    unk = json.loads(json.dumps(s))
    unk["grid"]["Group 5"]["Day 1"]["Period 1"] = "Kayaking"
    uf = os.path.join(t, "unknown.json")
    json.dump(unk, open(uf, "w", encoding="utf-8"))
    code, out = run("check", uf)
    ok(code == 1 and "'Kayaking' is not in the activities list" in out, "an unknown activity is caught, not guessed")

    code, out = run("table", sched)
    ok(code == 0 and out.count("| Group 18 |") == 5, "the paste-in table has one row per group for each of 5 days")

    def dump(obj, name):
        f = os.path.join(t, name)
        json.dump(obj, open(f, "w", encoding="utf-8"))
        return f

    # Andy 2: a fixed period that overfills an area makes draft itself exit with the check's code.
    tight = json.loads(json.dumps(SETUP))
    tight["rules"]["fixed"] = [{"groups": "all", "days": "all", "period": "Period 4", "activity": "Gaga"}]
    code, out = run("draft", dump(tight, "tight.json"), os.path.join(t, "tight-out.json"), "--seed", "1")
    ok(code == 1 and "Double booked: Gaga pit, Day 1, Period 4" in out,
       "red control: draft exits 1 when its own fixed periods double book an area")
    small = json.loads(json.dumps(SETUP))
    small.pop("rules")
    small["areas"] = [{"name": a["name"], "holds": 1} for a in SETUP["areas"]]
    code, out = run("draft", dump(small, "small.json"), os.path.join(t, "small-out.json"), "--seed", "1")
    ok(code == 0 and "Cannot fit: there are 18 groups but room for only 12" in out,
       "more groups than room gives a plain cannot-fit line")
    code, out = run("draft", setup, os.path.join(t, "again.json"), "--seed", "7")
    ok("Cannot fit" not in out, "control: no cannot-fit line when there is room")

    # Andy 3: a loosely spelled group or period is checked, never read as open.
    loose = json.loads(json.dumps(s))
    g1 = loose["grid"].pop("Group 1")
    g1["day 2"] = g1.pop("Day 2")
    g1["day 2"]["period 1 "] = "Gaga"
    del g1["day 2"]["Period 1"]
    loose["grid"]["group 1"] = g1
    loose["grid"]["Group 2"]["Day 2"]["Period 1"] = "Gaga"
    code, out = run("check", dump(loose, "loose.json"))
    ok(code == 1 and "Double booked: Gaga pit, Day 2, Period 1: Group 1, Group 2" in out,
       "red control: a double booking written with loose spelling is still caught")
    stray = json.loads(json.dumps(s))
    stray["grid"]["Group 19"] = stray["grid"]["Group 1"]
    code, out = run("check", dump(stray, "stray.json"))
    ok(code == 1 and "group called 'Group 19' that is not in the setup" in out, "an unknown group is caught")

    # Andy 4: a workbook whose labels or group rows changed is refused, in plain words.
    wb = load_workbook(book)
    wb["Master"].cell(row=2, column=2).value = "Monday"
    relabeled = os.path.join(t, "relabeled.xlsx")
    wb.save(relabeled)
    code, out = run("check", relabeled)
    ok(code == 2 and "day labels on the Master sheet were changed" in out, "red control: changed day labels are refused")
    wb = load_workbook(book)
    wb["Master"].cell(row=3, column=3).value = "Snack"
    periods = os.path.join(t, "periods.xlsx")
    wb.save(periods)
    code, out = run("check", periods)
    ok(code == 2 and "period labels on the Master sheet were changed" in out, "red control: changed period labels are refused")
    wb = load_workbook(book)
    wb["Master"].cell(row=4 + 18, column=1).value = "Group 19"
    extra = os.path.join(t, "extra.xlsx")
    wb.save(extra)
    code, out = run("check", extra)
    ok(code == 2 and "row for 'Group 19'" in out, "red control: an added group row is refused")

    # Andy 5 and Trevor 4: a name that looks like a formula is stored as text.
    formula = json.loads(json.dumps(plain))
    formula["groups"] = ["=SUM(A1)"] + GROUPS[1:]
    formula["title"] = "=1+1"
    fs = os.path.join(t, "formula-schedule.json")
    code, out = run("draft", dump(formula, "formula.json"), fs, "--seed", "2")
    fbook = os.path.join(t, "formula.xlsx")
    code, out = run("build", fs, fbook)
    ok(code == 0, "a group named like a formula still builds")
    wb = load_workbook(fbook) if os.path.exists(fbook) else None
    cells = [wb["Master"].cell(row=4, column=1), wb["Master"]["A1"]] if wb else []
    ok(bool(cells) and all(c.data_type == "s" for c in cells) and cells[0].value == "=SUM(A1)",
       "red control: '=SUM(A1)' and the title are stored as text, never as a formula")
    code, out = run("check", fbook)
    ok(code == 0, "the workbook with a formula-like name checks clean")
    ok(not [f for f in os.listdir(t) if f.endswith(".tmp.xlsx")], "no temporary workbook is left behind")

    # An activity limited to some groups: giving it to another group is reported.
    lim = json.loads(json.dumps(s))
    lim["grid"]["Group 1"]["Day 1"]["Period 1"] = "Archery"
    for g in GROUPS[1:]:
        if lim["grid"][g]["Day 1"]["Period 1"] == "Archery":
            lim["grid"][g]["Day 1"]["Period 1"] = ""
    code, out = run("check", dump(lim, "limited.json"))
    ok("Group 1 has Archery on Day 1, Period 1, but Archery is only for" in out,
       "red control: an activity given to a group it is not for is reported")

    print(f"{len(fails)} wrong")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
